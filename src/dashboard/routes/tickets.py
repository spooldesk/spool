"""Tickets list (with filters + pagination), per-ticket detail page, Excel export."""

from __future__ import annotations

import io
import logging
from datetime import datetime
from typing import Optional
from urllib.parse import quote, urlencode

from fastapi import APIRouter, File, Form, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, JSONResponse, RedirectResponse, StreamingResponse

from src.dashboard import queries
from src.dashboard.formatting import fmt_seconds as _fmt_seconds
from src.dashboard.i18n import t as tr

logger = logging.getLogger(__name__)

router = APIRouter()

# Cap on uploaded attachments forwarded to users (Telegram bots allow up to 50MB)
_MAX_UPLOAD_BYTES = 20 * 1024 * 1024  # 20 MB — sane default for support files
_MAX_UPLOAD_FILES = 10                # per reply; Telegram albums cap here too
# Uploads are held in memory to be checked before any of them is sent, so this
# is what one reply may cost the process, not just what the operator may attach.
_MAX_UPLOAD_TOTAL_BYTES = 100 * 1024 * 1024

PAGE_SIZE = 25


def _parse_date(value: str | None) -> Optional[datetime]:
    if not value:
        return None
    try:
        return datetime.strptime(value, "%Y-%m-%d")
    except ValueError:
        return None


# Anything wider than this cannot reach SQLite: binding it raises OverflowError
# from the driver, which surfaces as a 500 rather than as a rejected parameter.
#
# The bound is the column's, not the engine's. These filters are compared
# against `Integer` columns — admin ids, team ids — which PostgreSQL stores in
# 32 bits, so a value between 2^31 and 2^63 used to pass this check and be
# refused by the driver instead.
_MAX_ROW_ID = 2 ** 31 - 1


def _parse_int(value: str | None) -> Optional[int]:
    """Parse a query string value as int. Empty/invalid/out of range → None."""
    if not value:
        return None
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return None
    if not -_MAX_ROW_ID <= parsed <= _MAX_ROW_ID:
        return None
    return parsed


def _build_page_url(page: int, base_params: dict) -> str:
    """Build /tickets?<filters>&page=N URL — Python-side, no template macros."""
    params = {k: v for k, v in base_params.items() if v}
    params["page"] = str(page)
    return f"/tickets?{urlencode(params)}"


def _build_export_url(base_params: dict) -> str:
    """Build /tickets/export.xlsx?<filters> URL — same filters, no page."""
    params = {k: v for k, v in base_params.items() if v}
    if not params:
        return "/tickets/export.xlsx"
    return f"/tickets/export.xlsx?{urlencode(params)}"


@router.get("/tickets", response_class=HTMLResponse)
async def tickets_list(
    request: Request,
    # All query params as str so empty form values from <select> don't 422.
    status: str | None = None,
    team_id: str | None = None,
    category: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    q: str | None = None,
    assigned: str | None = None,
    stage: str | None = None,
    page: str | None = None,
):
    templates = request.app.state.templates

    # Normalise & validate
    status = "closed" if status == "closed" else "active"
    from src.services.workflow_service import STAGES
    if stage not in STAGES or status == "closed":
        stage = None
    team_id_int = _parse_int(team_id)
    if team_id_int is not None and team_id_int <= 0:
        team_id_int = None
    category = category or None
    df = _parse_date(date_from)
    dt = _parse_date(date_to)
    q = (q or "").strip()[:200] or None  # cap length; empty → no search
    page_int = max(0, _parse_int(page) or 0)

    # Operator filter: "me" → current operator's id, "none" → unassigned, "<id>" → that operator
    from src.dashboard.auth import current_user
    me = current_user(request)
    assigned_norm: int | str | None = None
    if assigned == "me" and me is not None:
        assigned_norm = me.id
    elif assigned == "none":
        assigned_norm = "none"
    elif assigned and _parse_int(assigned) is not None and assigned.isdigit():
        assigned_norm = int(assigned)
    else:
        assigned = ""  # normalize junk for links/template

    rows, total = await queries.list_tickets(
        status=status,
        team_id=team_id_int,
        category=category,
        date_from=df,
        date_to=dt,
        q=q,
        assigned=assigned_norm,
        stage=stage,
        page=page_int,
        page_size=PAGE_SIZE,
    )
    from src.services import assignment_service
    operators = [
        {"id": a.id, "label": a.name or a.email.split("@")[0]}
        for a in await assignment_service.list_active_operators()
    ]
    teams = await queries.list_teams_for_filter()
    categories = await queries.list_categories_for_filter()

    total_pages = (total + PAGE_SIZE - 1) // PAGE_SIZE if total else 1

    base_params = {
        "status": status or "",
        "team_id": str(team_id_int) if team_id_int else "",
        "category": category or "",
        "date_from": date_from or "",
        "date_to": date_to or "",
        "q": q or "",
        "assigned": assigned or "",
        "stage": stage or "",
    }

    prev_url = _build_page_url(page_int - 1, base_params) if page_int > 0 else None
    next_url = _build_page_url(page_int + 1, base_params) if page_int + 1 < total_pages else None
    export_url = _build_export_url(base_params)
    # Switching queues preserves search/filter context and starts at page one.
    # Stage describes live work and must never hide records in the archive.
    tab_params = {k: v for k, v in base_params.items() if v and k not in {"status", "stage"}}
    status_urls = {state: "/tickets?" + urlencode({**tab_params, "status": state})
                   for state in ("active", "closed")}
    status_counts = await queries.ticket_status_counts(
        team_id=team_id_int, category=category, date_from=df, date_to=dt,
        q=q, assigned=assigned_norm,
    )

    return templates.TemplateResponse(
        request,
        "tickets.html",
        {
            "active_page": "tickets",
            "rows": rows,
            "total": total,
            "page": page_int,
            "total_pages": total_pages,
            "page_size": PAGE_SIZE,
            "teams": teams,
            "categories": categories,
            "filters": {
                "status": status or "",
                "team_id": team_id_int or "",
                "category": category or "",
                "date_from": date_from or "",
                "date_to": date_to or "",
                "q": q or "",
                "assigned": assigned or "",
                "stage": stage or "",
            },
            "operators": operators,
            "prev_url": prev_url,
            "next_url": next_url,
            "export_url": export_url,
            "status_urls": status_urls,
            "status_counts": status_counts,
            "reset_url": "/tickets?" + urlencode({"status": status}),
            "saved_views": await _views_for(request, "/tickets"),
            "current_query": request.url.query,
            "fmt_seconds": _fmt_seconds,
        },
    )


@router.get("/board", response_class=HTMLResponse)
async def board_page(
    request: Request,
    team_id: str | None = None,
    assigned: str | None = None,
    q: str | None = None,
):
    """Kanban view of the same tickets the list shows, grouped by stage."""
    from src.dashboard.auth import current_user
    from src.services import assignment_service
    from src.services.workflow_service import BOARD_COLUMNS, STAGES

    templates = request.app.state.templates
    team_id_int = _parse_int(team_id)
    if team_id_int is not None and team_id_int <= 0:
        team_id_int = None
    q = (q or "").strip()[:200] or None

    me = current_user(request)
    assigned_norm: int | str | None = None
    if assigned == "me" and me is not None:
        assigned_norm = me.id
    elif assigned == "none":
        assigned_norm = "none"
    elif assigned and _parse_int(assigned) is not None and assigned.isdigit():
        assigned_norm = int(assigned)
    else:
        assigned = ""

    columns = await queries.board_columns(
        team_id=team_id_int, assigned=assigned_norm, q=q,
    )
    operators = [
        {"id": a.id, "label": a.name or a.email.split("@")[0]}
        for a in await assignment_service.list_active_operators()
    ]
    teams = await queries.list_teams_for_filter()

    return templates.TemplateResponse(
        request,
        "board.html",
        {
            "active_page": "board",
            "column_keys": list(BOARD_COLUMNS),
            "columns": columns,
            # For the card's move menu. Only the stages a card can be moved TO:
            # closed follows from status, not from a stage change.
            "stage_labels": {s: tr(f"stage_{s}") for s in STAGES},
            # Drives the take-next button and its counter.
            "queue_size": await queries.queue_size(me.id if me else None, "open"),
            "saved_views": await _views_for(request, "/board"),
            "current_query": request.url.query,
            "teams": teams,
            "operators": operators,
            "filters": {
                "team_id": team_id_int or "",
                "assigned": assigned or "",
                "q": q or "",
            },
            "fmt_seconds": _fmt_seconds,
        },
    )


@router.post("/api/tickets/{ticket_id}/stage")
async def api_set_stage(request: Request, ticket_id: int, stage: str = Form(...)):
    """Drag-and-drop endpoint for the board. Returns JSON so the card can be
    reverted in place when the move is refused (e.g. the ticket got closed)."""
    from src.services import workflow_service

    ok = await workflow_service.set_stage(ticket_id, (stage or "").strip())
    if not ok:
        return JSONResponse(
            {"ok": False, "error": tr("stage_locked")}, status_code=409,
        )
    await _auto_assign_to_me(request, ticket_id)
    return JSONResponse({"ok": True, "stage": stage})


@router.post("/api/tickets/{ticket_id}/presence", response_class=JSONResponse)
async def ticket_presence(
    request: Request,
    ticket_id: int,
    typing: str = Form(""),
    leaving: str = Form(""),
    audience: str = Form("public"),
    client_id: str = Form("", max_length=32),
):
    """Heartbeat: say I am here, learn who else is.

    Also nudges Telegram's own typing indicator so the customer sees that
    somebody is composing — that direction works; the reverse does not, since a
    bot is never told when a user is typing.
    """
    from src.services import presence_service as presence
    from src.dashboard.auth import current_user

    me = current_user(request)
    if me is None:
        return JSONResponse({"ok": False}, status_code=403)

    if client_id and (len(client_id) != 32 or any(c not in '0123456789abcdef' for c in client_id)):
        return JSONResponse({"ok": False}, status_code=400)
    name = me.name or me.email.split("@")[0]
    participant = f"operator:{me.id}" + (f":{client_id}" if client_id else "")
    if leaving == "1":
        await presence.heartbeat(ticket_id, participant, name, operator_id=me.id, leaving=True)
        return JSONResponse({"ok": True, "others": []})

    is_typing = typing == "1" and getattr(request.state, 'section_permissions', {}).get('tickets') == 'edit'
    others = await presence.heartbeat(ticket_id, participant, name,
                                     operator_id=me.id, typing=is_typing,
                                     public=audience == "public")
    if is_typing and audience == "public":
        await _show_typing_to_customer(ticket_id)

    return JSONResponse({"ok": True, "others": others})


# Telegram shows "typing…" for about five seconds per call, so repeating any
# faster only spends requests.
_TYPING_MIN_INTERVAL_SEC = 4.0
_last_typing_sent: dict[int, float] = {}


async def _show_typing_to_customer(ticket_id: int) -> None:
    """Best-effort: never let a decoration break the heartbeat."""
    import time as _time

    now = _time.monotonic()
    if now - _last_typing_sent.get(ticket_id, 0.0) < _TYPING_MIN_INTERVAL_SEC:
        return
    _last_typing_sent[ticket_id] = now

    try:
        from sqlalchemy import select as _select
        from sqlalchemy.orm import selectinload as _selectinload

        from src.dashboard.bot_client import get_bot
        from src.database.db import db_session
        from src.database.models import Ticket

        async with db_session() as session:
            ticket = (await session.execute(
                _select(Ticket).where(Ticket.id == ticket_id, Ticket.status == "active")
                .options(_selectinload(Ticket.user))
            )).scalar_one_or_none()
        if ticket is None or ticket.user is None:
            return
        from src.channels import get as _get_channel
        from src.channels.identity import address_of as _addr, channel_of as _chan
        await _get_channel(_chan(ticket.user)).typing(_addr(ticket.user))
    except Exception as e:  # noqa: BLE001
        logger.debug(f"typing action skipped for ticket #{ticket_id}: {e}")


@router.get("/api/queue-status")
async def queue_status():
    """Sidebar badge + new-ticket notification poll (auth via router deps)."""
    return await queries.queue_status()


@router.get("/api/tickets-version")
async def tickets_version():
    """Asked before a live refresh: unchanged token means skip the re-render."""
    return {"version": await queries.tickets_version()}


@router.get("/tickets/export.xlsx")
async def tickets_export_xlsx(
    request: Request,
    status: str | None = None,
    team_id: str | None = None,
    category: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
    q: str | None = None,
    assigned: str | None = None,
    stage: str | None = None,
):
    """Excel export of tickets matching the same filters as /tickets.

    A custom_reason from a sensitive category is masked here exactly as it
    is in the dashboard — an export is not a way around that.
    Bodies of messages are NOT included — only headers / counts / metadata.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from src.services.workflow_service import STAGES

    if status not in ("active", "closed"):
        status = None
    if stage not in STAGES:
        stage = None
    team_id_int = _parse_int(team_id)
    if team_id_int is not None and team_id_int <= 0:
        team_id_int = None
    category = category or None
    df = _parse_date(date_from)
    dt = _parse_date(date_to)
    q = (q or "").strip()[:200] or None

    from src.dashboard.auth import current_user
    me = current_user(request)
    assigned_norm: int | str | None = None
    if assigned == "me" and me is not None:
        assigned_norm = me.id
    elif assigned == "none":
        assigned_norm = "none"
    elif assigned and assigned.isdigit() and _parse_int(assigned) is not None:
        assigned_norm = int(assigned)

    rows = await queries.list_tickets_for_export(
        status=status,
        team_id=team_id_int,
        category=category,
        date_from=df,
        date_to=dt,
        q=q,
        assigned=assigned_norm,
        stage=stage,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = tr("x_tickets")

    headers = [
        "ID", tr("x_status"), tr("x_stage"), tr("x_operator"), tr("x_team"),
        tr("x_username"), tr("x_tg_id"), tr("x_local_id"),
        tr("x_category"), tr("x_created_at"), tr("x_closed_at"),
        tr("x_resolution_h"), tr("x_resolution_m"),
        tr("x_msgs_total"), tr("x_msgs_user"), tr("x_msgs_admin"), tr("x_msgs_template"),
        tr("x_close_reasons"), tr("x_custom_reason"), tr("x_csat"),
    ]
    ws.append(headers)

    # Header styling
    header_fill = PatternFill("solid", fgColor="1F2937")  # slate-800
    header_font = Font(bold=True, color="FFFFFF")
    header_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = header_align
    ws.row_dimensions[1].height = 28

    for r in rows:
        sec = r["resolution_seconds"]
        hours = round(sec / 3600, 2) if sec is not None else None
        minutes = round(sec / 60, 1) if sec is not None else None

        ws.append([
            r["id"],
            r["status"],
            tr("stage_" + (r["stage"] or "new")) if r["status"] == "active" else "",
            r["operator"],
            r["team"],
            r["username"],
            r["tg_id"],
            r["local_user_id"],
            r["category"],
            r["created_at"].strftime("%Y-%m-%d %H:%M") if r["created_at"] else "",
            r["closed_at"].strftime("%Y-%m-%d %H:%M") if r["closed_at"] else "",
            hours,
            minutes,
            r["msg_total"],
            r["msg_user"],
            r["msg_admin"],
            r["via_template"],
            r["closure_categories"],
            r["custom_reason"],
            r["csat_rating"] if r["csat_rating"] else "",
        ])

    # Column widths — picked by eye, good enough without measuring text
    widths = [8, 10, 13, 16, 16, 22, 14, 12, 22, 18, 18, 14, 14, 12, 10, 12, 14, 28, 36, 8]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Freeze header row + auto-filter so users can sort/filter in Excel
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

    # Stream the workbook back without writing a temp file
    buf = io.BytesIO()
    from src.dashboard.exporting import literal_strings
    literal_strings(wb)
    wb.save(buf)
    buf.seek(0)

    # Filename includes the active filters and current date — easier to file later
    parts = [datetime.utcnow().strftime("%Y-%m-%d")]
    if status:
        parts.append(status)
    if team_id_int:
        parts.append(f"team{team_id_int}")
    if df or dt:
        parts.append(f"{date_from or 'X'}_to_{date_to or 'X'}")
    filename = "tickets_" + "_".join(parts) + ".xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


from src.channels import caps as _channel_caps  # noqa: E402


@router.get("/tickets/{ticket_id}", response_class=HTMLResponse)
async def ticket_detail_page(request: Request, ticket_id: int,
                             sent: str | None = None, err: str | None = None,
                             queue: str | None = None):
    templates = request.app.state.templates
    detail = await queries.ticket_detail(ticket_id)
    if detail is None:
        raise HTTPException(status_code=404, detail=tr("e_ticket_not_found"))

    # Closure reasons + ticket categories from the editable catalogs (same the
    # bot uses). Reason "value" is the snapshot name stored on the ticket.
    from src.database.db import db_session
    from src.services import catalog_service, close_reason_service
    async with db_session() as session:
        close_categories = [
            {"value": r.name, "label": close_reason_service.display_name(r)}
            for r in await close_reason_service.list_active(session)
        ]
        ticket_categories = [
            {"name": c.name, "emoji": c.emoji or ""}
            for c in await catalog_service.list_active(session)
        ]
    close_categories.append({
        "value": close_reason_service.CUSTOM_VALUE,
        "label": close_reason_service.custom_display(),
    })

    from src.services import assignment_service
    operators = [
        {"id": a.id, "label": a.name or a.email.split("@")[0]}
        for a in await assignment_service.list_active_operators(ticket_id)
    ]

    return templates.TemplateResponse(
        request,
        "ticket_detail.html",
        {
            "active_page": "tickets",
            "t": detail,
            "fmt_seconds": _fmt_seconds,
            "flash_sent": sent,
            "flash_err": err,
            "close_categories": close_categories,
            "ticket_categories": ticket_categories,
            "operators": operators,
            # The composer enforces the same limits the route does, so an
            # over-sized pick is refused before the upload, not after it.
            # "queue=<source>" means keep going; "queue=empty" means the last
            # one is answered and there is nothing left to hand over.
            "queue_source": queue if queue in QUEUE_SOURCES else None,
            "queue_empty": queue == "empty",
            "queue_size": await queries.queue_size(
                _current_operator_id(request), queue if queue in QUEUE_SOURCES else "open"),
            "max_upload_files": _MAX_UPLOAD_FILES,
            "max_upload_bytes": _MAX_UPLOAD_BYTES,
            # Whether the composer offers the attach button. The input itself
            # is always rendered — the script reads it on load, and making the
            # element conditional once killed every behaviour registered after
            # that line. See tests/test_composer_regressions.py.
            "channel_takes_files": _channel_caps(detail.get("channel")).files,
        },
    )


@router.post("/tickets/{ticket_id}/assign")
async def ticket_assign(ticket_id: int, admin_id: str = Form("")):
    """Assign the ticket to an operator ('' clears the assignment)."""
    from urllib.parse import quote
    from src.services import assignment_service

    target = int(admin_id) if admin_id.isdigit() else None
    ok = await assignment_service.assign_ticket(ticket_id, target)
    if not ok:
        return RedirectResponse(
            f"/tickets/{ticket_id}?err={quote(tr("e_assign_failed"))}",
            status_code=303)
    return RedirectResponse(f"/tickets/{ticket_id}?sent=assigned", status_code=303)


@router.post("/tickets/{ticket_id}/stage")
async def ticket_set_stage(ticket_id: int, stage: str = Form("")):
    """Move an open ticket between board stages (new / in_progress / waiting)."""
    from urllib.parse import quote
    from src.services import workflow_service

    ok = await workflow_service.set_stage(ticket_id, (stage or "").strip())
    if not ok:
        return RedirectResponse(
            f"/tickets/{ticket_id}?err={quote(tr('stage_locked'))}", status_code=303)
    return RedirectResponse(f"/tickets/{ticket_id}?sent=stage", status_code=303)


@router.post("/tickets/{ticket_id}/note")
async def ticket_add_note(request: Request, ticket_id: int, text: str = Form("")):
    """Add an internal note. This never reaches the user — no Bot API call here,
    on purpose: the note lives in the dashboard timeline only."""
    from urllib.parse import quote
    from src.services import workflow_service

    note_id = await workflow_service.add_internal_note(
        ticket_id, text=text, operator_id=_current_operator_id(request))
    if note_id is None:
        return RedirectResponse(
            f"/tickets/{ticket_id}?err={quote(tr('note_empty'))}", status_code=303)
    await _auto_assign_to_me(request, ticket_id)
    return RedirectResponse(f"/tickets/{ticket_id}?sent=note", status_code=303)


# A double-click on Send used to reach the customer twice. Disabling the button
# covers the common case, but not a refresh or a Back-then-resubmit, so the
# server refuses an identical message to the same ticket inside this window.
# Short on purpose: sending the same words twice within seconds is a stutter,
# sending them a minute apart is a person repeating themselves deliberately.
_DUPLICATE_WINDOW_SEC = 8


async def _is_duplicate_send(ticket_id: int, text: str, reply_to_id: int | None = None) -> bool:
    """True if this exact text just went to this ticket from an operator."""
    from sqlalchemy import select as _select

    from src.database.db import db_session
    from src.database.models import Message as MessageRecord

    text = (text or "").strip()
    if not text:
        return False
    async with db_session() as session:
        last = (await session.execute(
            _select(MessageRecord)
            .where(MessageRecord.ticket_id == ticket_id,
                   MessageRecord.sender_type == "admin",
                   MessageRecord.is_internal == False)  # noqa: E712
            .order_by(MessageRecord.id.desc())
            .limit(1)
        )).scalar_one_or_none()
    if (last is None or (last.text or "").strip() != text
            or last.reply_to_id != reply_to_id):
        return False
    age = (datetime.utcnow() - last.created_at).total_seconds() if last.created_at else None
    return age is not None and 0 <= age < _DUPLICATE_WINDOW_SEC


@router.post("/tickets/{ticket_id}/message/{message_id}/delete")
async def ticket_delete_message(request: Request, ticket_id: int, message_id: int):
    """Withdraw a delivered message from the customer's chat."""
    from urllib.parse import quote

    from src.dashboard.bot_client import get_bot
    from src.services import message_service

    from src.database.db import db_session
    from src.database.models import Message as MessageRecord
    async with db_session() as session:
        message = await session.get(MessageRecord, message_id)
        if message is None or message.ticket_id != ticket_id:
            raise HTTPException(status_code=404, detail=tr("e_ticket_not_found"))
        from src.dashboard.auth import current_user
        actor = current_user(request)
        if actor and actor.role not in {"owner", "admin"} and message.operator_id != actor.id:
            raise HTTPException(status_code=403, detail=tr("chat_edit_forbidden"))

    result = await message_service.delete_for_user(get_bot(), message_id, actor_id=actor.id if actor else None)
    if not result.ok:
        return RedirectResponse(
            f"/tickets/{ticket_id}?err={quote(tr(result.error_key))}", status_code=303)
    return RedirectResponse(f"/tickets/{ticket_id}?sent=deleted", status_code=303)


@router.post("/tickets/{ticket_id}/message/{message_id}/edit")
async def ticket_edit_message(request: Request, ticket_id: int, message_id: int,
                              text: str = Form(...), revision: int = Form(...)):
    from src.dashboard.auth import current_user
    from src.services.conversation_service import edit_message
    actor = current_user(request)
    if actor is None:
        raise HTTPException(status_code=403)
    result = await edit_message(ticket_id, message_id, actor.id, text, revision)
    return JSONResponse({"ok": result.ok, "error": tr(result.error_key) if not result.ok else ""},
                        status_code=200 if result.ok else 409)


@router.post('/tickets/{ticket_id}/message/{message_id}/reconcile')
async def reconcile_message(request: Request, ticket_id: int, message_id: int,
                            revision: int = Form(...), applied: str = Form(...)):
    from src.dashboard.auth import current_user
    from src.services.conversation_service import reconcile_mutation
    actor = current_user(request)
    if actor is None or applied not in {'yes','no'}:
        raise HTTPException(status_code=403)
    result = await reconcile_mutation(ticket_id,message_id,actor.id,revision,applied=='yes')
    return JSONResponse({'ok':result.ok,'error':tr(result.error_key) if not result.ok else ''},
                        status_code=200 if result.ok else 409)


@router.post("/tickets/{ticket_id}/message/{message_id}/retry")
async def ticket_retry_message(request: Request, ticket_id: int, message_id: int,
                               acknowledge_duplicate: str = Form("")):
    """Retry only this ticket's intent; ambiguous sends require explicit consent."""
    from sqlalchemy import select
    from src.dashboard.auth import require_role
    from src.database.db import db_session
    from src.database.delivery_models import OutboundJob
    from src.database.models import Message as MessageRecord, Ticket
    from src.services import outbound

    await require_role("owner", "admin", "agent")(request)
    async with db_session() as session:
        message = await session.get(MessageRecord, message_id)
        ticket = await session.get(Ticket, ticket_id)
        if (message is None or message.ticket_id != ticket_id or ticket is None
                or message.sender_type != "admin" or message.is_internal
                or message.deleted_at):
            raise HTTPException(status_code=404, detail=tr("e_ticket_not_found"))
        job = (await session.execute(select(OutboundJob).where(
            OutboundJob.message_id == message_id,
        ))).scalar_one_or_none()
        if job is None:
            raise HTTPException(status_code=404, detail=tr("e_attachment_not_found"))
        if job.status == "uncertain" and acknowledge_duplicate != "yes":
            raise HTTPException(status_code=409, detail=tr("delivery_duplicate_risk"))
        job_id = job.id
    result = await outbound.retry(job_id, allow_uncertain=acknowledge_duplicate == "yes")
    suffix = "sent=1" if result.ok else "err=" + quote(tr(result.error_key or "e_send_failed"))
    return RedirectResponse(f"/tickets/{ticket_id}?{suffix}", status_code=303)



def _standing_queue(request: Request) -> str:
    """"Always hand me the next one", from the operator's preferences.

    The `queue` form field means "keep going" for this one click, which is
    what the Send-and-next button sends. Someone who set the preference is
    asking for that click to be implied on every reply, so it stands in
    where the field is empty — and only there, so an ordinary Send from a
    ticket opened out of the blue still leaves them where they were.
    """
    prefs = getattr(request.state, "prefs", None) or {}
    return "open" if prefs.get("queue_advance") else ""

@router.post("/tickets/{ticket_id}/reply")
async def ticket_reply(
    request: Request,
    ticket_id: int,
    text: str = Form(""),
    # Not list[UploadFile]: an empty file input still posts a part, and one
    # with no filename arrives as a plain string, which that type rejects with
    # a 422 — breaking the ordinary text-only reply.
    attachments: list[UploadFile | str] = File(default=[]),
    reply_to_id: str = Form(""),
    # Set when the operator arrived from the queue: sending should hand
    # them the next one, not the conversation they just finished.
    queue: str = Form(""),
):
    """Send an admin reply (text and/or a file) to the ticket's user from web.

    Every reply is recorded as a durable intent before contacting the channel.
    """
    from urllib.parse import quote
    from src.config import load_settings
    from src.channels import get as get_channel
    from src.channels.identity import address_of, channel_of
    from src.dashboard.bot_client import get_bot
    from src.database.db import db_session
    from src.services import outbound, template_service

    text = (text or "").strip()
    # The operator's own sign-off, if they set one. Applied before the
    # duplicate check rather than after, so two identical replies are still
    # two identical replies once signed.
    prefs = getattr(request.state, "prefs", None) or {}
    if text and prefs.get("signature"):
        from src.services.preferences_service import apply_signature
        text = apply_signature(text, prefs["signature"])
    # Duck-typed on purpose: through that union pydantic hands back Starlette's
    # UploadFile, not FastAPI's subclass, so isinstance against the imported
    # name quietly matches nothing and every upload falls into the text branch.
    files = [f for f in (attachments or []) if getattr(f, "filename", None)]
    has_file = bool(files)
    parent_id = None
    if reply_to_id:
        from src.services.conversation_service import reply_context
        try:
            parent_id = int(reply_to_id)
            if not 0 < parent_id <= 2147483647:
                raise ValueError("Invalid reply")
            async with db_session() as session:
                await reply_context(session, ticket_id, parent_id)
        except ValueError:
            return RedirectResponse(f"/tickets/{ticket_id}?err={quote(tr('chat_reply_unavailable'))}", status_code=303)

    # Silently succeed rather than error: the operator meant to send this once,
    # and it was sent once. Telling them off for a stutter would be noise.
    if not has_file and await _is_duplicate_send(ticket_id, text, parent_id):
        logger.info(f"Ignored duplicate reply to ticket #{ticket_id}")
        return RedirectResponse(f"/tickets/{ticket_id}?sent=1", status_code=303)
    if not text and not has_file:
        return RedirectResponse(f"/tickets/{ticket_id}?err={quote(tr("e_empty_message"))}", status_code=303)

    settings = load_settings()
    bot = get_bot()

    # ----- File branch -----
    if has_file:
        from sqlalchemy import select
        from sqlalchemy.orm import selectinload

        from src.database.models import Ticket

        if len(files) > _MAX_UPLOAD_FILES:
            return RedirectResponse(
                f"/tickets/{ticket_id}?err={quote(tr('e_too_many_files', limit=_MAX_UPLOAD_FILES))}",
                status_code=303)

        # Read and check everything before sending anything. Rejecting the
        # fourth file after three have already reached the customer is not a
        # failure the operator can undo.
        payloads: list[tuple[str, bytes]] = []
        total = 0
        for f in files:
            content = await f.read(_MAX_UPLOAD_BYTES + 1)
            if len(content) > _MAX_UPLOAD_BYTES:
                return RedirectResponse(
                    f"/tickets/{ticket_id}?err={quote(tr('e_file_too_large_x', name=f.filename))}",
                    status_code=303)
            total += len(content)
            if total > _MAX_UPLOAD_TOTAL_BYTES:
                return RedirectResponse(
                    f"/tickets/{ticket_id}?err={quote(tr('e_files_total_too_large'))}",
                    status_code=303)
            payloads.append((f.filename, content))

        sent_count = 0
        failure: str | None = None
        try:
            async with db_session() as session:
                ticket = (await session.execute(
                    select(Ticket).where(Ticket.id == ticket_id)
                    .options(selectinload(Ticket.user))
                )).scalar_one_or_none()
                if ticket is None or ticket.status != "active":
                    return RedirectResponse(
                        f"/tickets/{ticket_id}?err={quote(tr("e_ticket_not_active"))}", status_code=303)
                if ticket.user is None:
                    return RedirectResponse(
                        f"/tickets/{ticket_id}?err={quote(tr("e_no_recipient"))}", status_code=303)
                channel = get_channel(channel_of(ticket.user))
                address = address_of(ticket.user)
                # Asked before anything is read back to the operator, not after
                # three files have gone nowhere.
                if not channel.caps.files:
                    return RedirectResponse(
                        f"/tickets/{ticket_id}?err={quote(tr("e_no_files_on_channel"))}",
                        status_code=303)

            for name, content in payloads:
                # The text is a caption on the first file only — repeating it
                # under every attachment would read like five separate replies.
                caption = text if (text and sent_count == 0) else None
                async with db_session(write=True) as session:
                    job = await outbound.enqueue(
                        session, ticket_id=ticket_id, channel=channel.name, address=address,
                        filename=name, content=content, text=caption or "",
                        operator_id=_current_operator_id(request),
                        reply_to_id=parent_id,
                        require_active=True,
                    )
                    job_id, message_id = job.id, job.message_id
                delivery = await outbound.deliver(job_id)
                if not delivery.ok:
                    # Whatever already left is real and stays recorded below.
                    failure = tr(delivery.error_key or "e_send_failed")
                    break
                # The mirror uses the durable row and records its forum id.
                from src.services.announce_service import announce_message
                try:
                    await announce_message(bot, ticket_id, message_id)
                except Exception:
                    logger.warning("Could not mirror delivered attachment %s", message_id)

                sent_count += 1
        except Exception as e:
            return RedirectResponse(
                f"/tickets/{ticket_id}?err={quote(tr("e_file_error", what=type(e).__name__))}",
                status_code=303)

        if sent_count:
            await _auto_assign_to_me(request, ticket_id)
            nxt = await _queue_redirect(
                request, queue or _standing_queue(request), ticket_id)
            if nxt and not failure:
                return nxt
        if failure:
            # Say how far it got: "nothing was sent" and "three of five went"
            # need very different things from the operator.
            logger.warning(
                f"Reply to ticket #{ticket_id}: {sent_count}/{len(payloads)} files sent, then {failure}")
            return RedirectResponse(
                f"/tickets/{ticket_id}?err="
                f"{quote(tr('e_files_partial', done=sent_count, total=len(payloads), what=failure))}",
                status_code=303)
        return RedirectResponse(f"/tickets/{ticket_id}?sent=1", status_code=303)

    # ----- Text-only branch -----
    try:
        async with db_session() as session:
            result = await template_service.apply_text(
                session, bot,
                ticket_id=ticket_id,
                text=text,
                admin_tg_id=0,
                operator_id=_current_operator_id(request),
                admin_chat_id=settings.admin_chat_id,
                source_label=tr("src_web_dashboard"),
                reply_to_id=parent_id,
            )
            await session.commit()
    except Exception as e:
        return RedirectResponse(
            f"/tickets/{ticket_id}?err={quote(tr("e_generic", what=type(e).__name__))}", status_code=303)
    if not result.success:
        return RedirectResponse(
            f"/tickets/{ticket_id}?err={quote(result.error or tr("e_send_failed"))}", status_code=303)
    # (stage advance happens inside template_service.apply_text)
    await _auto_assign_to_me(request, ticket_id)
    # The text-only reply is the common one, so the loop has to continue from
    # here too — wiring only the attachment branch left the queue working for
    # the rarer path and silently stopping on the usual one.
    nxt = await _queue_redirect(request, queue or _standing_queue(request), ticket_id)
    if nxt:
        return nxt
    return RedirectResponse(f"/tickets/{ticket_id}?sent=1", status_code=303)


async def _auto_assign_to_me(request: Request, ticket_id: int) -> None:
    """First responder from the dashboard claims an unassigned ticket."""
    from src.dashboard.auth import current_user
    from src.services import assignment_service
    me = current_user(request)
    if me is not None:
        try:
            await assignment_service.auto_assign(ticket_id, me.id)
        except Exception:  # noqa: BLE001 — assignment must never break a reply
            pass



async def _queue_redirect(request: Request, queue: str, just_finished: int):
    """Where to go after a reply when the operator is working the queue."""
    if queue not in QUEUE_SOURCES:
        return None
    ticket_id = await _claim_next(request, queue, after=just_finished)
    if ticket_id is None:
        return RedirectResponse(f"/tickets/{just_finished}?queue=empty", status_code=303)
    return RedirectResponse(f"/tickets/{ticket_id}?queue={queue}", status_code=303)


def _current_operator_id(request: Request) -> int | None:
    from src.dashboard.auth import current_user
    me = current_user(request)
    return me.id if me is not None else None


# ==================== Saved views ====================


@router.post("/views/save")
async def saved_view_create(request: Request, name: str = Form(""),
                            path: str = Form("/tickets"), query: str = Form("")):
    """Name the filters currently applied."""
    from src.services import saved_view_service

    operator_id = _current_operator_id(request)
    if operator_id is None:
        return RedirectResponse(path, status_code=303)

    view = await saved_view_service.save(
        operator_id, name=name, path=path, query=query,
    )
    target = path if path in saved_view_service.ALLOWED_PATHS else "/tickets"
    if view is None:
        return RedirectResponse(f"{target}?err={quote(tr('sv_not_saved'))}",
                                status_code=303)
    return RedirectResponse(f"{target}?{view.query}" if view.query else target,
                            status_code=303)


@router.post("/views/{view_id}/delete")
async def saved_view_delete(request: Request, view_id: int,
                            path: str = Form("/tickets")):
    from src.services import saved_view_service

    operator_id = _current_operator_id(request)
    if operator_id is not None:
        await saved_view_service.remove(operator_id, view_id)
    return RedirectResponse(
        path if path in saved_view_service.ALLOWED_PATHS else "/tickets",
        status_code=303,
    )


async def _views_for(request: Request, path: str) -> list:
    """This operator's saved views for one page, or nothing when signed out."""
    from src.services import saved_view_service

    operator_id = _current_operator_id(request)
    if operator_id is None:
        return []
    return await saved_view_service.list_for(operator_id, path)


# ==================== The work queue ====================

QUEUE_SOURCES = ("open", "mine", "unassigned", "all")


async def _claim_next(request: Request, source: str,
                      after: int | None = None) -> int | None:
    """Pick the next ticket owed an answer and claim it for this operator.

    Claiming is the point. A queue that only suggests hands the same ticket to
    two people, which is the collision it was supposed to remove — so the pick
    and the assignment happen together, before anyone is sent anywhere.
    """
    from src.dashboard.auth import current_user
    from src.services import assignment_service

    me = current_user(request)
    operator_id = me.id if me is not None else None

    ticket_id = await queries.queue_next(operator_id, source, exclude=after)
    if ticket_id is None or operator_id is None:
        return ticket_id
    try:
        await assignment_service.auto_assign(ticket_id, operator_id)
    except Exception:  # noqa: BLE001 — never block the operator over bookkeeping
        logger.warning(f"Could not claim #{ticket_id} for operator {operator_id}")
    return ticket_id


@router.post("/queue/next")
async def queue_take_next(request: Request, source: str = Form("open")):
    """POST, not GET: this assigns a ticket, and a link that mutates can be
    fired by a browser prefetching it."""
    if source not in QUEUE_SOURCES:
        source = "open"
    ticket_id = await _claim_next(request, source)
    if ticket_id is None:
        return RedirectResponse("/board?queue=empty", status_code=303)
    return RedirectResponse(f"/tickets/{ticket_id}?queue={source}", status_code=303)

@router.post("/tickets/{ticket_id}/close")
async def ticket_close(
    request: Request,
    ticket_id: int,
    categories: list[str] = Form(default=[]),
    custom_reason: str = Form(""),
):
    """Close a ticket from web with reason categories + optional custom text."""
    from urllib.parse import quote
    from datetime import datetime as _dt
    from sqlalchemy import select
    from sqlalchemy.orm import selectinload
    from src.config import load_settings
    from src.dashboard.bot_client import get_bot
    from src.database.db import db_session
    from src.database.models import Message as MessageRecord, Ticket

    settings = load_settings()
    bot = get_bot()
    custom_reason = (custom_reason or "").strip()
    categories = [c for c in (categories or []) if c]

    from src.utils.i18n import get_text, CLOSURE_SYSTEM_TEXT

    user_tg_id = None
    user_lang = None
    from src.services import csat_service
    try:
        async with db_session(write=True) as session:
            ticket = (await session.execute(
                select(Ticket).where(Ticket.id == ticket_id, Ticket.status == "active")
                .options(selectinload(Ticket.user))
            )).scalar_one_or_none()
            if ticket is None:
                return RedirectResponse(
                    f"/tickets/{ticket_id}?err={quote(tr("e_ticket_not_found_or_closed"))}",
                    status_code=303)
            user_tg_id = ticket.user.tg_id if ticket.user else None
            user_lang = ticket.user.language if ticket.user else None
            from src.services.ticket_closure import close_in_session
            job = await close_in_session(session, ticket, ticket.user,
                operator_id=_current_operator_id(request), categories=categories,
                custom_reason=custom_reason, reason_prefix=tr('msg_reason_prefix'))
            closing_job_id = job.id
            await session.commit()
    except Exception as e:
        return RedirectResponse(
            f"/tickets/{ticket_id}?err={quote(tr("e_generic", what=type(e).__name__))}", status_code=303)

    if user_tg_id:
        # Closure notice + CSAT rating request (never raises)
        from src.services import csat_service
        await csat_service.send_close_notification(
            bot, user_tg_id=user_tg_id, ticket_id=ticket_id, lang=user_lang, job_id=closing_job_id,
        )
    return RedirectResponse(f"/tickets/{ticket_id}?sent=closed", status_code=303)


@router.post("/tickets/{ticket_id}/category")
async def ticket_set_category(ticket_id: int, category: str = Form("")):
    """Assign/replace the ticket's category from the card (admin-driven).

    Used when the bot didn't ask the user up front, or to correct a category.
    Stores the name as a snapshot string on the ticket.
    """
    from urllib.parse import quote
    from sqlalchemy import select
    from src.database.db import db_session
    from src.database.models import Ticket

    category = (category or "").strip()
    try:
        async with db_session() as session:
            ticket = (await session.execute(
                select(Ticket).where(Ticket.id == ticket_id)
            )).scalar_one_or_none()
            if ticket is None:
                return RedirectResponse(
                    f"/tickets/{ticket_id}?err={quote(tr("e_ticket_not_found"))}", status_code=303)
            ticket.category = category
            await session.commit()
    except Exception as e:
        return RedirectResponse(
            f"/tickets/{ticket_id}?err={quote(tr("e_generic", what=type(e).__name__))}", status_code=303)
    return RedirectResponse(f"/tickets/{ticket_id}?sent=category", status_code=303)


@router.post("/tickets/{ticket_id}/ai-suggest", response_class=JSONResponse)
async def ticket_ai_suggest(ticket_id: int):
    """Draft from published knowledge; operator reviews before sending."""
    try:
        from sqlalchemy import select
        from src.database.db import db_session
        from src.database.models import Message
        from src.ai.knowledge import suggest_from_knowledge
        async with db_session() as session:
            question = (await session.execute(select(Message.text).where(
                Message.ticket_id == ticket_id, Message.sender_type == "user",
                Message.deleted_at.is_(None),
            ).order_by(Message.id.desc()).limit(1))).scalar_one_or_none()
        result = await suggest_from_knowledge(ticket_id, question or "")
        if result.suggested_text:
            return JSONResponse({"ok": True, "draft": result.suggested_text,
                                 "sources": result.sources, "ai_call_id": result.ai_call_id})
        return JSONResponse({"ok": False, "error": tr("e_ai_no_answer")})
    except Exception:
        logger.warning("Knowledge draft unavailable for ticket %s", ticket_id)
        return JSONResponse({"ok": False, "error": tr("e_ai_no_answer")})


@router.get("/tickets/attachment/{msg_id}")
async def ticket_attachment(msg_id: int):
    """Stream an attachment through the dashboard, from whichever channel it
    arrived on.

    Only a reference is stored, never the bytes — a Telegram file_id, a Slack
    private URL — so fetching means asking the channel. That used to be a
    direct Bot API call here, which made every attachment a Telegram
    attachment. Sensitive-category tickets are refused outright.
    """
    import io
    from sqlalchemy import select
    from sqlalchemy.orm import selectinload
    from src.channels import get as get_channel
    from src.channels.identity import channel_of
    from src.dashboard import queries as _q
    from src.dashboard.routes.widget import _attachment_headers, _content_type, _safe_filename
    from src.database.db import db_session
    from src.database.models import Message as MessageRecord, Ticket
    from src.database.delivery_models import OutboundJob
    from src.services import file_store

    async with db_session() as session:
        m = (await session.execute(
            select(MessageRecord)
            .where(MessageRecord.id == msg_id)
            .options(selectinload(MessageRecord.ticket).selectinload(Ticket.user))
        )).scalar_one_or_none()
        if m is None or not m.file_id:
            raise HTTPException(status_code=404, detail=tr("e_attachment_not_found"))
        category = m.ticket.category if m.ticket else None
        if _q.is_sensitive_category(category):
            raise HTTPException(status_code=403, detail=tr("e_hidden_sensitive"))
        author = m.ticket.user if m.ticket else None
        channel = channel_of(author) if author is not None else None
        file_id = m.file_id
        ftype = (m.file_type or "").lower()
        file_name = _safe_filename(m.file_name or (
            "photo.jpg" if "photo" in ftype else "attachment"
        ))
        # MIME metadata arrives from customers on several channels. Only
        # raster images may render on the dashboard's authenticated origin.
        mime = _content_type(file_name)
        stored_key = (await session.execute(select(OutboundJob.file_key).where(
            OutboundJob.message_id == m.id,
        ))).scalar_one_or_none()

    try:
        data = file_store.get(stored_key) if stored_key else None
        if data is None:
            data = await get_channel(channel).fetch_file(file_id)
    except LookupError:
        # The customer is on a channel this process cannot reach — a Slack
        # ticket on an instance whose credentials were removed, say.
        data = None
    if not data:
        raise HTTPException(status_code=502, detail=tr("e_attachment_unavailable"))

    return StreamingResponse(
        io.BytesIO(data),
        media_type=mime,
        headers=_attachment_headers(file_name),
    )


@router.get("/tickets/{ticket_id}/templates.json", response_class=JSONResponse)
async def ticket_templates_json(ticket_id: int):
    """Reply templates for this ticket's team, with placeholders rendered.

    Returns ready-to-send text so the admin can drop it into the composer and
    tweak before sending.
    """
    from sqlalchemy import select
    from sqlalchemy.orm import selectinload
    from src.database.db import db_session
    from src.database.models import Team, Ticket
    from src.services import template_service

    async with db_session() as session:
        ticket = (await session.execute(
            select(Ticket).where(Ticket.id == ticket_id).options(selectinload(Ticket.user))
        )).scalar_one_or_none()
        if ticket is None:
            return JSONResponse({"ok": False, "error": tr("e_ticket_not_found")}, status_code=404)

        templates = await template_service.list_for_ticket(session, ticket, limit=50)
        team = None
        if ticket.team_id:
            team = (await session.execute(
                select(Team).where(Team.id == ticket.team_id)
            )).scalar_one_or_none()
        ctx = template_service.build_context(ticket, ticket.user, team, admin_first_name="")
        items = [
            {"id": t.id, "title": t.title, "category": t.category or "",
             "body": template_service.render(t.body, ctx)}
            for t in templates
        ]
    return JSONResponse({"ok": True, "templates": items})


@router.get("/tickets/{ticket_id}/messages.json", response_class=JSONResponse)
async def ticket_messages_json(request: Request, ticket_id: int, after_id: str | None = None):
    """What the open page has not seen yet.

    `after_id` is the client's cursor. Without it the whole conversation comes
    back, which is what a page that has just loaded actually wants; with it the
    answer to a quiet ticket is a handful of bytes.
    """
    cursor = _parse_int(after_id) or 0
    payload = await queries.messages_after(ticket_id, cursor)
    if payload is None:
        return JSONResponse({"ok": False}, status_code=404)
    from src.dashboard.auth import current_user
    actor = current_user(request)
    if actor is not None and getattr(request.state, 'section_permissions', {}).get('tickets') != 'edit':
        for delivery in payload.get("deliveries", []):
            delivery["can_retry"] = False
    return JSONResponse({"ok": True, **payload})
