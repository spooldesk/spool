"""Managerial reports: period summary with deltas, operators, teams, AI impact.

RBAC: admin + owner (wired in app.py). The Excel export mirrors the page —
one file a support lead can forward upstairs.
"""

from __future__ import annotations

import io
import logging
from datetime import date, datetime, time, timedelta

from fastapi import APIRouter, Request
from fastapi.responses import HTMLResponse, StreamingResponse

from src.dashboard import report_queries
from src.dashboard.i18n import t as tr

logger = logging.getLogger(__name__)

router = APIRouter()

_PRESETS = ("7d", "30d", "this_week", "last_week", "this_month", "last_month")


def _period_bounds(preset: str | None, date_from: str | None, date_to: str | None,
                   today: date | None = None) -> tuple[datetime, datetime, str]:
    """Resolve (from, to_exclusive, effective_preset). Custom dates win."""
    today = today or datetime.utcnow().date()

    def _parse(raw: str | None) -> date | None:
        try:
            return datetime.strptime(raw, "%Y-%m-%d").date() if raw else None
        except ValueError:
            return None

    df, dt = _parse(date_from), _parse(date_to)
    if df and dt and df <= dt:
        # Inclusive UI dates → exclusive upper bound
        return (datetime.combine(df, time.min),
                datetime.combine(dt + timedelta(days=1), time.min), "custom")

    preset = preset if preset in _PRESETS else "7d"
    monday = today - timedelta(days=today.weekday())
    first_of_month = today.replace(day=1)

    if preset == "7d":
        start, end = today - timedelta(days=6), today + timedelta(days=1)
    elif preset == "30d":
        start, end = today - timedelta(days=29), today + timedelta(days=1)
    elif preset == "this_week":
        start, end = monday, today + timedelta(days=1)
    elif preset == "last_week":
        start, end = monday - timedelta(days=7), monday
    elif preset == "this_month":
        start, end = first_of_month, today + timedelta(days=1)
    else:  # last_month
        last_month_end = first_of_month
        start = (first_of_month - timedelta(days=1)).replace(day=1)
        end = last_month_end

    return datetime.combine(start, time.min), datetime.combine(end, time.min), preset


async def _gather(date_from: datetime, date_to: datetime) -> dict:
    """Current period + the immediately preceding period of equal length."""
    span = date_to - date_from
    prev_from, prev_to = date_from - span, date_from
    return {
        "summary": await report_queries.report_summary(date_from, date_to),
        "prev": await report_queries.report_summary(prev_from, prev_to),
        "operators": await report_queries.report_by_operator(date_from, date_to),
        "teams": await report_queries.report_by_team(date_from, date_to),
    }


def _delta_pct(cur, prev) -> float | None:
    """% change vs previous period; None when it can't be computed honestly."""
    if cur is None or prev is None or prev == 0:
        return None
    return (cur - prev) / prev * 100


def _fmt_minutes(m: float | None) -> str:
    if m is None:
        return "—"
    if m < 90:
        return f"{m:.0f} {tr('min_short')}"
    return f"{m / 60:.1f} {tr('hour_short')}"

def _fmt_hours(h: float | None) -> str:
    if h is None:
        return "—"
    if h < 1:
        return f"{h * 60:.0f} {tr('min_short')}"
    if h < 24:
        return f"{h:.1f} {tr('hour_short')}"
    return f"{h / 24:.1f} {tr('day_short')}"


def _build_cards(s: dict, p: dict) -> list[dict]:
    """Summary cards: value + Δ to the previous equal period.

    `invert` marks metrics where growth is BAD (times) so the template can
    colour the delta correctly.
    """
    return [
        {"label": tr("r_created"), "value": str(s["created"]),
         "delta": _delta_pct(s["created"], p["created"]), "invert": False},
        {"label": tr("r_closed"), "value": str(s["closed"]),
         "delta": _delta_pct(s["closed"], p["closed"]), "invert": False},
        {"label": tr("r_first_response"), "value": _fmt_minutes(s["avg_first_response_min"]),
         "delta": _delta_pct(s["avg_first_response_min"], p["avg_first_response_min"]),
         "invert": True},
        {"label": tr("r_close_time"), "value": _fmt_hours(s["avg_close_hours"]),
         "delta": _delta_pct(s["avg_close_hours"], p["avg_close_hours"]), "invert": True},
        {"label": "CSAT",
         "value": (f"{s['csat_avg']:.2f} / 5" if s["csat_avg"] is not None else "—"),
         "hint": (f"{s['csat_count']} {tr('ratings')}" if s["csat_count"]
                  else tr("kpi_csat_hint_none")),
         "delta": _delta_pct(s["csat_avg"], p["csat_avg"]), "invert": False},
        {"label": tr("r_ai_closed"), "value": str(s["ai_closed"]),
         "hint": (tr("r_ai_closed_pct", pct=f"{s['ai_closed_pct']:.0f}")
                  if s["ai_closed_pct"] is not None else None),
         "delta": _delta_pct(s["ai_closed"], p["ai_closed"]), "invert": False},
    ]


@router.get("/reports", response_class=HTMLResponse)
async def reports_page(
    request: Request,
    preset: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
):
    templates = request.app.state.templates
    df, dt, effective = _period_bounds(preset, date_from, date_to)
    data = await _gather(df, dt)

    from urllib.parse import urlencode
    export_qs = urlencode({
        "date_from": df.date().isoformat(),
        "date_to": (dt.date() - timedelta(days=1)).isoformat(),
    })

    return templates.TemplateResponse(request, "reports.html", {
        "active_page": "reports",
        "preset": effective,
        "presets": [(key, tr(f"preset_{key}")) for key in _PRESETS],
        "date_from": df.date().isoformat(),
        # Show the inclusive end date in the UI (exclusive bound minus a day)
        "date_to": (dt.date() - timedelta(days=1)).isoformat(),
        "cards": _build_cards(data["summary"], data["prev"]),
        "export_url": f"/reports/export.xlsx?{export_qs}",
        "fmt_hours": _fmt_hours,
        **data,
    })


@router.get("/reports/export.xlsx")
async def reports_export(
    preset: str | None = None,
    date_from: str | None = None,
    date_to: str | None = None,
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    df, dt, _ = _period_bounds(preset, date_from, date_to)
    data = await _gather(df, dt)
    s, p = data["summary"], data["prev"]

    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="FFFFFF")

    def _style_header(ws, ncols: int):
        for col in range(1, ncols + 1):
            c = ws.cell(row=1, column=col)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 24

    def _fmt(v, digits=1):
        return round(v, digits) if isinstance(v, float) else (v if v is not None else "")

    # --- Sheet 1: summary with previous-period comparison
    ws = wb.active
    ws.title = tr("x_summary")
    ws.append([tr("x_metric"), tr("x_period"), tr("x_prev_period")])
    _style_header(ws, 3)
    period_label = f"{df.date().isoformat()} — {(dt.date() - timedelta(days=1)).isoformat()}"
    rows = [
        (tr("x_dates"), period_label, ""),
        (tr("x_created"), s["created"], p["created"]),
        (tr("x_closed"), s["closed"], p["closed"]),
        (tr("x_first_response_min"), _fmt(s["avg_first_response_min"]), _fmt(p["avg_first_response_min"])),
        (tr("x_close_hours"), _fmt(s["avg_close_hours"]), _fmt(p["avg_close_hours"])),
        (tr("x_csat_avg"), _fmt(s["csat_avg"], 2), _fmt(p["csat_avg"], 2)),
        (tr("x_csat_count"), s["csat_count"], p["csat_count"]),
        (tr("x_ai_closed"), s["ai_closed"], p["ai_closed"]),
        (tr("x_ai_share"), _fmt(s["ai_closed_pct"]), _fmt(p["ai_closed_pct"])),
        (tr("x_ai_offered"), s["deflect_offered"], p["deflect_offered"]),
        (tr("x_ai_accepted"), s["deflect_accepted"], p["deflect_accepted"]),
    ]
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate([36, 22, 22], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # --- Sheet 2: operators
    ws2 = wb.create_sheet(tr("x_operators"))
    ws2.append([tr("x_operator"), tr("th_closed_cnt"), tr("x_csat"), tr("x_ratings"),
                tr("x_avg_close_h"), tr("x_in_progress")])
    _style_header(ws2, 6)
    for op in data["operators"]:
        ws2.append([
            op["label"], op["closed"], _fmt(op["csat_avg"], 2), op["csat_count"],
            _fmt(op["avg_close_hours"]), op["active_now"],
        ])
    for i, w in enumerate([26, 10, 8, 8, 20, 16], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # --- Sheet 3: teams
    ws3 = wb.create_sheet(tr("x_teams"))
    ws3.append([tr("x_team"), tr("x_created"), tr("x_closed"), tr("x_csat")])
    _style_header(ws3, 4)
    for t in data["teams"]:
        ws3.append([t["name"], t["created"], t["closed"], _fmt(t["csat_avg"], 2)])
    for i, w in enumerate([26, 10, 10, 8], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"report_{df.date().isoformat()}_{(dt.date() - timedelta(days=1)).isoformat()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
