"""Release gates: revocation, concurrent owner protection and literal exports."""

import asyncio
import io
import threading
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, update

from src.dashboard import auth
from src.dashboard.app import app, TEMPLATES
from src.dashboard.i18n import strings
from src.database.db import db_session
from src.database.models import AdminUser, Team, Ticket, User
from src.services import admin_user_service as admins
from tests.conftest import next_tg_id


@pytest.fixture()
async def operator(initialized_db, monkeypatch):
    monkeypatch.setitem(TEMPLATES.env.globals, "T", strings("en"))
    return await admins.create_admin(
        email=f"release-{next_tg_id()}@example.test", password="release-password",
        role="owner", name="=1+2",
    )


def _client(operator):
    c = TestClient(app, base_url="http://testserver")
    c.cookies.set(auth.SESSION_COOKIE,
                  auth.make_session_cookie(operator.id, operator.session_version))
    c.get("/login", follow_redirects=False)
    return c


async def test_password_reset_revokes_every_old_cookie_and_login_recovers(operator):
    c = _client(operator)
    old = c.cookies.get(auth.SESSION_COOKIE)
    assert c.get("/tickets").status_code == 200
    assert await admins.set_password(operator.id, "replacement-password")
    assert not await auth.is_current_session(old)
    assert c.get("/tickets", follow_redirects=False).status_code == 307
    assert c.get("/login", follow_redirects=False).status_code == 200
    r = c.post("/login", data={
        "email": operator.email, "password": "replacement-password",
        "csrf_token": c.cookies.get("supbot_csrf"),
    }, follow_redirects=False)
    assert r.status_code == 303
    assert c.get("/tickets").status_code == 200
    c.close()


@pytest.mark.parametrize("mutation", ["disable", "role", "delete"])
async def test_operator_changes_revoke_old_cookies(operator, mutation):
    cookie = auth.make_session_cookie(operator.id, operator.session_version)
    if mutation == "disable":
        await admins.set_active(operator.id, False)
    elif mutation == "role":
        await admins.update_profile(operator.id, role="agent")
    else:
        await admins.delete_admin(operator.id)
    assert not await auth.is_current_session(cookie)


async def test_logout_revokes_stolen_copy_and_stale_logout_cannot_revoke_new_login(operator):
    c = _client(operator)
    old = c.cookies.get(auth.SESSION_COOKIE)
    assert c.post("/logout", follow_redirects=False).status_code == 403
    assert await auth.is_current_session(old)
    r = c.post("/logout", data={"csrf_token": c.cookies.get("supbot_csrf")},
               follow_redirects=False)
    assert r.status_code == 303
    assert not await auth.is_current_session(old)
    latest = await admins.authenticate(operator.email, "release-password")
    new = auth.make_session_cookie(latest.id, latest.session_version)
    assert await auth.is_current_session(new)
    assert not await admins.revoke_sessions(operator.id, expected_version=operator.session_version)
    assert await auth.is_current_session(new)
    c.close()


async def test_password_change_during_login_cannot_issue_a_fresh_cookie_for_old_password(
    operator, monkeypatch,
):
    started, release = threading.Event(), threading.Event()
    real_verify = admins._verify

    def delayed_verify(plain, hashed):
        result = real_verify(plain, hashed)
        started.set()
        assert release.wait(timeout=10)
        return result

    monkeypatch.setattr(admins, "_verify", delayed_verify)
    task = asyncio.create_task(admins.authenticate(operator.email, "release-password"))
    try:
        assert await asyncio.to_thread(started.wait, 10)
        await admins.set_password(operator.id, "replacement-password")
    finally:
        release.set()
    assert await task is None


def test_legacy_uid_only_cookie_is_rejected():
    legacy = auth._signer().sign(b"42").decode()
    assert auth.session_claims(legacy) is None


@pytest.mark.parametrize("action", ["demote", "disable", "delete"])
async def test_concurrent_mutations_cannot_remove_both_remaining_owners(initialized_db, action):
    first = await admins.create_admin(email=f"owners-a-{next_tg_id()}@example.test",
                                      password="release-password", role="owner")
    second = await admins.create_admin(email=f"owners-b-{next_tg_id()}@example.test",
                                       password="release-password", role="owner")
    async with db_session(write=True) as session:
        previous = list((await session.execute(select(
            AdminUser.id, AdminUser.role, AdminUser.is_active,
        ).where(AdminUser.id.not_in([first.id, second.id])))).all())
        # Isolate the invariant from unrelated accounts in this shared test DB.
        await session.execute(update(AdminUser).where(
            AdminUser.id.not_in([first.id, second.id]), AdminUser.role == "owner",
        ).values(role="admin"))

    async def remove(uid):
        if action == "demote":
            return await admins.update_profile(uid, role="agent")
        if action == "disable":
            return await admins.set_active(uid, False)
        return await admins.delete_admin(uid)

    try:
        results = await asyncio.gather(remove(first.id), remove(second.id),
                                       return_exceptions=True)
        assert results.count(True) == 1
        assert sum(isinstance(r, admins.LastOwnerError) for r in results) == 1
        async with db_session() as session:
            remaining = (await session.execute(select(AdminUser.id).where(
                AdminUser.is_active.is_(True), AdminUser.role == "owner",
            ))).scalars().all()
        assert len(remaining) == 1
    finally:
        async with db_session(write=True) as session:
            for uid, role, active in previous:
                await session.execute(update(AdminUser).where(AdminUser.id == uid).values(
                    role=role, is_active=active))
            await session.execute(update(AdminUser).where(
                AdminUser.id.in_([first.id, second.id]),
            ).values(role="agent", is_active=False))


async def test_ticket_export_preserves_assignment_filter_and_literal_text(operator):
    async with db_session() as session:
        team = Team(name="=1+3")
        session.add(team)
        await session.flush()
        ids = []
        for assigned in (operator.id, None):
            user = User(tg_id=next_tg_id(), username="=1+1", team_id=team.id)
            session.add(user)
            await session.flush()
            ticket = Ticket(user_id=user.id, team_id=team.id, category="=1+4",
                            status="active", assigned_admin_id=assigned)
            session.add(ticket)
            await session.flush()
            ids.append(ticket.id)
        team_id = team.id
    c = _client(operator)
    for assigned, expected in (("me", ids[0]), (str(operator.id), ids[0]), ("none", ids[1])):
        response = c.get("/tickets/export.xlsx", params={"team_id": team_id, "assigned": assigned})
        assert response.status_code == 200
        sheet = load_workbook(io.BytesIO(response.content), data_only=False).active
        assert [row[0] for row in sheet.iter_rows(min_row=2, values_only=True)] == [expected]
        strings_found = [cell for row in sheet.iter_rows() for cell in row
                         if isinstance(cell.value, str) and cell.value.startswith("=")]
        assert strings_found
        assert all(cell.data_type == "s" for cell in strings_found)
    c.close()


async def test_report_export_never_infers_formulas_from_names(operator):
    async with db_session() as session:
        team = Team(name="=2+3")
        session.add(team)
        await session.flush()
        user = User(tg_id=next_tg_id(), team_id=team.id)
        session.add(user)
        await session.flush()
        session.add(Ticket(user_id=user.id, team_id=team.id,
                           category="Network",
                           assigned_admin_id=operator.id, status="closed",
                           created_at=datetime.utcnow(), closed_at=datetime.utcnow()))
    c = _client(operator)
    response = c.get("/reports/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content), data_only=False)
    matching = [cell for sheet in workbook for row in sheet for cell in row
                if cell.value in ("=1+2", "=2+3")]
    assert {cell.value for cell in matching} == {"=1+2", "=2+3"}
    assert all(cell.data_type == "s" for cell in matching)
    c.close()
